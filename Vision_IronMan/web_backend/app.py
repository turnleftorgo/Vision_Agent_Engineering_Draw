from __future__ import annotations

import json
import re
import shutil
import threading
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pydantic import BaseModel
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .orchestrator import process_pdf_run, run_log
from .store import ScanStore


PROJECT_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_DIR / "web_data"
RUNS_DIR = DATA_DIR / "runs"
BACKEND_INSTANCE_ID = uuid.uuid4().hex
STORE = ScanStore(DATA_DIR / "claw_view.sqlite3", BACKEND_INSTANCE_ID)

app = FastAPI(title="Claw View local scan API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://127.0.0.1:4173", "http://localhost:4173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def reject_runs_from_previous_backend(request, call_next):
    path_parts = request.url.path.strip("/").split("/")
    if len(path_parts) >= 3 and path_parts[:2] == ["api", "runs"]:
        run_id = path_parts[2]
        run = STORE.get_run(run_id)
        if run and run.get("backend_instance_id") != BACKEND_INSTANCE_ID:
            return JSONResponse(
                status_code=410,
                content={"detail": "该任务属于已停止的后端进程，无法恢复"},
            )
    return await call_next(request)


class CropUpdate(BaseModel):
    bbox: list[int]


class RecordFieldUpdate(BaseModel):
    field: str
    value: str | None


class MetadataUpdate(BaseModel):
    field: str
    value: str | None


def public_record(record: dict[str, Any]) -> dict[str, Any]:
    value = dict(record)
    run_id = value["run_id"]
    value["crop_url"] = f"/api/runs/{run_id}/records/{value['id']}/crop"
    value["page_url"] = f"/api/runs/{run_id}/pages/{value['page_index']}"
    for key in ("model_crop_path", "user_crop_path", "effective_crop_path", "source_page_path"):
        value.pop(key, None)
    return value


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok", "backend_instance_id": BACKEND_INSTANCE_ID}


@app.post("/api/runs")
def create_run(pdf: UploadFile = File(...)) -> dict[str, Any]:
    if not pdf.filename or not pdf.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="请选择 PDF 文件")
    run_id = uuid.uuid4().hex
    run_dir = RUNS_DIR / run_id
    run_dir.mkdir(parents=True, exist_ok=False)
    pdf_path = run_dir / "source.pdf"
    with pdf_path.open("wb") as destination:
        shutil.copyfileobj(pdf.file, destination)
    try:
        if pdf_path.read_bytes()[:5] != b"%PDF-":
            raise ValueError
    except ValueError:
        shutil.rmtree(run_dir)
        raise HTTPException(status_code=400, detail="上传文件不是有效 PDF")
    STORE.create_run(run_id, pdf.filename, pdf_path, run_dir)
    threading.Thread(
        target=process_pdf_run,
        args=(STORE, run_id),
        name=f"scan-{run_id[:8]}",
        daemon=True,
    ).start()
    run_log(run_id, f"frontend uploaded PDF: {pdf.filename}", component="API")
    return {
        "run_id": run_id,
        "status": "queued",
        "backend_instance_id": BACKEND_INSTANCE_ID,
    }


@app.get("/api/runs/{run_id}")
def get_run(run_id: str) -> dict[str, Any]:
    run = STORE.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="扫描任务不存在")
    run.pop("pdf_path", None)
    run.pop("run_dir", None)
    run["records"] = [public_record(item) for item in STORE.list_records(run_id)]
    progress = STORE.get_progress(run_id)
    for page in progress["pages"]:
        page_index = page["page_index"]
        page["overview_url"] = f"/api/runs/{run_id}/progress/{page_index}/overview"
        for module_index, stage1_module in enumerate(page["stage1_modules"], start=1):
            stage1_module["crop_url"] = (
                f"/api/runs/{run_id}/progress/{page_index}/modules/"
                f"{module_index}/source"
            )
        for module in page["stage2_modules"].values():
            module["visualization_url"] = (
                f"/api/runs/{run_id}/progress/{page_index}/modules/"
                f"{module['module_index']}"
            )
            module.pop("visualization_path", None)
    run["progress"] = progress
    return run


@app.patch("/api/runs/{run_id}/records/{record_id}")
def update_record(run_id: str, record_id: str, request: RecordFieldUpdate) -> dict[str, Any]:
    if not STORE.get_record(run_id, record_id):
        raise HTTPException(status_code=404, detail="检验记录不存在")
    try:
        record = STORE.update_record_fields(run_id, record_id, {request.field: request.value})
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return public_record(record) if record else {}


@app.patch("/api/runs/{run_id}/metadata")
def update_metadata(run_id: str, request: MetadataUpdate) -> dict[str, str]:
    if not STORE.get_run(run_id):
        raise HTTPException(status_code=404, detail="扫描任务不存在")
    try:
        STORE.update_run_metadata(run_id, request.field, request.value)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return {"status": "ok"}


@app.get("/api/runs/{run_id}/export.xlsx")
def export_run_xlsx(run_id: str) -> StreamingResponse:
    run = STORE.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="扫描任务不存在")
    records = STORE.list_records(run_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inspection"
    headers = [
        "No.", "Project", "Revision", "Author", "Date", "Module", "Page",
        "FAI", "SPC", "Description", "Nominal", "USL", "LSL", "100%",
        "DC", "Points", "SPC截图",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="A8C7E4")
    thin = Side(style="thin", color="222222")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [7, 18, 14, 18, 13, 24, 9, 9, 12, 34, 14, 10, 10, 10, 10, 22, 38]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for number, record in enumerate(records, start=1):
        row = number + 1
        sheet.append([
            number, record.get("project"), record.get("revision"), record.get("author"),
            record.get("drawing_date"), record.get("module"), f"page {record['page_index']}",
            record.get("fai"), record.get("spc"), record.get("description"),
            record.get("nominal"), record.get("usl"), record.get("lsl"),
            record.get("hundred_percent"), record.get("dc"), record.get("points"), None,
        ])
        sheet.row_dimensions[row].height = 52
        for cell in sheet[row]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        image_path = Path(record["effective_crop_path"])
        if image_path.is_file():
            image = ExcelImage(str(image_path))
            scale = min(250 / image.width, 48 / image.height)
            image.width *= scale
            image.height *= scale
            sheet.add_image(image, f"Q{row}")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    base_name = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(run["pdf_name"]).stem).strip("._") or "inspection"
    filename = f"{base_name}_inspection.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}; filename*=UTF-8''{quote(filename)}"},
    )


@app.get("/api/runs/{run_id}/progress/{page_index}/modules/{module_index}/source")
def get_stage1_module_crop(run_id: str, page_index: int, module_index: int) -> FileResponse:
    """Return the Stage 1 module crop even before Stage 2 FAI detection finishes."""
    run = STORE.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="扫描任务不存在")
    with STORE.connect() as db:
        row = db.execute(
            """SELECT stage1_modules FROM page_progress
            WHERE run_id = ? AND page_index = ?""",
            (run_id, page_index),
        ).fetchone()
    modules = json.loads(row[0]) if row else []
    if module_index < 1 or module_index > len(modules):
        raise HTTPException(status_code=404, detail="模块不存在")
    crop_path = modules[module_index - 1].get("crop_path")
    if not crop_path:
        raise HTTPException(status_code=404, detail="模块截图尚未生成")
    path = (Path(run["run_dir"]) / "pipeline" / f"page_{page_index:04d}" / crop_path).resolve()
    if Path(run["run_dir"]).resolve() not in path.parents or not path.is_file():
        raise HTTPException(status_code=404, detail="模块截图不存在")
    return FileResponse(path, media_type="image/png")


@app.get("/api/runs/{run_id}/pages/{page_index}")
def get_page(run_id: str, page_index: int) -> FileResponse:
    run = STORE.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="扫描任务不存在")
    pages = sorted(
        (Path(run["run_dir"]) / "pages").glob("page-*.png"),
        key=lambda path: int(path.stem.rsplit("-", 1)[-1]),
    )
    if page_index < 1 or page_index > len(pages):
        raise HTTPException(status_code=404, detail="PDF 页面不存在")
    return FileResponse(pages[page_index - 1], media_type="image/png")


@app.get("/api/runs/{run_id}/records/{record_id}/crop")
def get_crop(run_id: str, record_id: str) -> FileResponse:
    record = STORE.get_record(run_id, record_id)
    if not record:
        raise HTTPException(status_code=404, detail="检验记录不存在")
    path = Path(record["effective_crop_path"])
    if not path.is_file():
        raise HTTPException(status_code=404, detail="截图文件不存在")
    return FileResponse(path, media_type="image/png")


@app.get("/api/runs/{run_id}/progress/{page_index}/overview")
def get_overview(run_id: str, page_index: int) -> FileResponse:
    run = STORE.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="扫描任务不存在")
    with STORE.connect() as db:
        row = db.execute(
            """SELECT overview_path FROM page_progress
            WHERE run_id = ? AND page_index = ?""",
            (run_id, page_index),
        ).fetchone()
    if not row or not row[0] or not Path(row[0]).is_file():
        raise HTTPException(status_code=404, detail="Stage 1 总览尚未生成")
    path = Path(row[0]).resolve()
    if Path(run["run_dir"]).resolve() not in path.parents:
        raise HTTPException(status_code=403, detail="无权访问该文件")
    return FileResponse(path, media_type="image/jpeg")


@app.get("/api/runs/{run_id}/progress/{page_index}/modules/{module_index}")
def get_module_visualization(run_id: str, page_index: int, module_index: int) -> FileResponse:
    run = STORE.get_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="扫描任务不存在")
    with STORE.connect() as db:
        row = db.execute(
            """SELECT stage2_modules FROM page_progress
            WHERE run_id = ? AND page_index = ?""",
            (run_id, page_index),
        ).fetchone()
    modules = json.loads(row[0]) if row else {}
    module = modules.get(str(module_index))
    path = Path(module["visualization_path"]).resolve() if module and module.get("visualization_path") else None
    if not path or not path.is_file():
        raise HTTPException(status_code=404, detail="模块 FAI 总览尚未生成")
    if Path(run["run_dir"]).resolve() not in path.parents:
        raise HTTPException(status_code=403, detail="无权访问该文件")
    return FileResponse(path, media_type="image/jpeg")


@app.put("/api/runs/{run_id}/records/{record_id}/crop")
def update_crop(run_id: str, record_id: str, request: CropUpdate) -> dict[str, Any]:
    record = STORE.get_record(run_id, record_id)
    if not record:
        raise HTTPException(status_code=404, detail="检验记录不存在")
    if len(request.bbox) != 4:
        raise HTTPException(status_code=422, detail="bbox 必须有四个整数")
    x1, y1, x2, y2 = request.bbox
    width, height = int(record["page_width"]), int(record["page_height"])
    x1, x2 = sorted((max(0, min(width, x1)), max(0, min(width, x2))))
    y1, y2 = sorted((max(0, min(height, y1)), max(0, min(height, y2))))
    if x2 - x1 < 8 or y2 - y1 < 8:
        raise HTTPException(status_code=422, detail="截图区域过小")

    source_path = Path(record["source_page_path"])
    run = STORE.get_run(run_id)
    if not run or not source_path.is_file():
        raise HTTPException(status_code=404, detail="原始 PDF 页面不存在")
    revision_dir = Path(run["run_dir"]) / "user_crops"
    revision_dir.mkdir(parents=True, exist_ok=True)
    existing = len(list(revision_dir.glob(f"{record_id}-*.png")))
    crop_path = revision_dir / f"{record_id}-{existing + 1:03d}.png"
    with Image.open(source_path) as image:
        image.convert("RGB").crop((x1, y1, x2, y2)).save(crop_path)
    STORE.save_user_crop(record_id, [x1, y1, x2, y2], crop_path)
    updated = STORE.get_record(run_id, record_id)
    run_log(
        run_id,
        f"frontend saved crop override: record={record_id}, bbox={[x1, y1, x2, y2]}",
        component="API",
    )
    return public_record(updated) if updated else {}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        "web_backend.app:app",
        host="127.0.0.1",
        port=8002,
        reload=True,
        access_log=False,
    )
